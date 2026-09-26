from dataclasses import dataclass, field
from typing import List, Optional, Dict
from pathlib import Path
import logging

logger = logging.getLogger("sage.rag.chunker")

@dataclass
class Chunk:
    content: str
    chunk_index: int
    page_number: Optional[int] = None
    section: Optional[str] = None
    char_offset: int = 0
    token_count: int = 0
    metadata: dict = field(default_factory=dict)

class DocumentChunker:
    def __init__(self, chunk_size: int = 512, chunk_overlap: int = 64):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap

    def chunk_file(self, file_path: str, file_type: str) -> List[Chunk]:
        """Route to type-specific chunker."""
        path = Path(file_path)
        if file_type == "pdf":
            return self._chunk_pdf(path)
        elif file_type == "docx":
            return self._chunk_docx(path)
        elif file_type in ("xlsx", "xls"):
            return self._chunk_xlsx(path)
        elif file_type in ("pptx",):
            return self._chunk_pptx(path)
        elif file_type in ("txt", "md", "text"):
            return self._chunk_text(path)
        else:
            return self._chunk_text(path)  # fallback: treat as plain text

    def _chunk_pdf(self, path: Path) -> List[Chunk]:
        import fitz
        chunks: List[Chunk] = []
        chunk_index = 0
        char_offset = 0
        
        try:
            with fitz.open(str(path)) as doc:
                for i, page in enumerate(doc):
                    text = page.get_text().strip()
                    if not text:
                        continue
                    
                    page_chunks = self._split_text(text, char_offset, page=i + 1, section=f"Page {i + 1}")
                    for sc in page_chunks:
                        sc.chunk_index = chunk_index
                        chunks.append(sc)
                        chunk_index += 1
                    
                    char_offset += len(text)
        except Exception as e:
            logger.error(f"Error chunking PDF {path}: {e}")
        return chunks

    def _chunk_docx(self, path: Path) -> List[Chunk]:
        import docx
        chunks: List[Chunk] = []
        try:
            doc = docx.Document(str(path))
            full_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
            if full_text:
                chunks = self._split_text(full_text, 0, page=None, section=None)
                for i, c in enumerate(chunks):
                    c.chunk_index = i
        except Exception as e:
            logger.error(f"Error chunking DOCX {path}: {e}")
        return chunks

    def _chunk_xlsx(self, path: Path) -> List[Chunk]:
        import openpyxl
        chunks: List[Chunk] = []
        chunk_index = 0
        try:
            wb = openpyxl.load_workbook(str(path), data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_texts = [str(cell) for cell in row if cell is not None]
                    if row_texts:
                        rows.append(" | ".join(row_texts))
                
                if rows:
                    text = "\n".join(rows)
                    sheet_chunks = self._split_text(text, 0, page=None, section=sheet_name)
                    for sc in sheet_chunks:
                        sc.chunk_index = chunk_index
                        sc.metadata["sheet"] = sheet_name
                        chunks.append(sc)
                        chunk_index += 1
        except Exception as e:
            logger.error(f"Error chunking XLSX {path}: {e}")
        return chunks

    def _chunk_pptx(self, path: Path) -> List[Chunk]:
        """
        Chunk a PPTX file by slides.
        Each slide becomes one or more chunks (split if content exceeds chunk_size).
        Metadata includes slide_number and slide_title.
        """
        from pptx import Presentation
        
        try:
            prs = Presentation(str(path))
        except Exception as e:
            logger.error(f"Error opening PPTX {path}: {e}")
            return []

        chunks: List[Chunk] = []
        chunk_index = 0
        char_offset = 0
        
        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()
            
            slide_texts: List[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)
                
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = " | ".join(
                            cell.text.strip() for cell in row.cells if cell.text.strip()
                        )
                        if row_text:
                            slide_texts.append(row_text)
            
            if not slide_texts:
                continue
            
            full_slide_text = "\n".join(slide_texts)
            
            slide_chunks = self._split_text(
                text=full_slide_text,
                base_offset=char_offset,
                page=slide_num,
                section=slide_title or f"Slide {slide_num}",
            )
            
            for sc in slide_chunks:
                sc.chunk_index = chunk_index
                sc.metadata["slide_number"] = slide_num
                if slide_title:
                    sc.metadata["slide_title"] = slide_title
                chunks.append(sc)
                chunk_index += 1
            
            char_offset += len(full_slide_text)
        
        return chunks

    def _chunk_text(self, path: Path) -> List[Chunk]:
        chunks: List[Chunk] = []
        try:
            with open(path, "r", encoding="utf-8") as f:
                text = f.read()
            if text.strip():
                chunks = self._split_text(text, 0, page=None, section=None)
                for i, c in enumerate(chunks):
                    c.chunk_index = i
        except Exception as e:
            logger.error(f"Error chunking TXT {path}: {e}")
        return chunks

    def _split_text(self, text: str, base_offset: int, page: Optional[int], section: Optional[str]) -> List[Chunk]:
        """Core splitting: character-based with overlap as a proxy for tokens."""
        # Note: in a real token-based chunker, we'd use a tokenizer. 
        # Using ~4 chars per token for this basic implementation.
        char_chunk_size = self.chunk_size * 4
        char_overlap = self.chunk_overlap * 4
        
        if not text:
            return []
            
        chunks = []
        start = 0
        text_len = len(text)
        
        while start < text_len:
            end = min(start + char_chunk_size, text_len)
            
            # If not at the end, try to break at a newline or space
            if end < text_len:
                last_newline = text.rfind('\n', start, end)
                if last_newline != -1 and last_newline > start + char_chunk_size // 2:
                    end = last_newline
                else:
                    last_space = text.rfind(' ', start, end)
                    if last_space != -1 and last_space > start + char_chunk_size // 2:
                        end = last_space
                        
            chunk_content = text[start:end].strip()
            if chunk_content:
                chunks.append(Chunk(
                    content=chunk_content,
                    chunk_index=0, # set by caller
                    page_number=page,
                    section=section,
                    char_offset=base_offset + start,
                    token_count=len(chunk_content) // 4,
                    metadata={}
                ))
                
            start = end - char_overlap
            
        return chunks
